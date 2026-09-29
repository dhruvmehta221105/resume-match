import re
import os
from pathlib import Path
from typing import Dict

import fitz  # PyMuPDF
from docx import Document


def parse_resume(file_path: str) -> Dict:
    """Parse a resume file (PDF or DOCX) and return structured data."""
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext == ".pdf":
        raw_text = _extract_pdf_text(file_path)
    elif ext in (".docx", ".doc"):
        raw_text = _extract_docx_text(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    raw_text = _clean_text(raw_text)

    return {
        "raw_text": raw_text,
        "candidate_name": _extract_name(raw_text),
        "email": _extract_email(raw_text),
        "phone": _extract_phone(raw_text),
        "word_count": len(raw_text.split()),
    }


def _extract_pdf_text(file_path: str) -> str:
    text_parts = []
    with fitz.open(file_path) as doc:
        for page in doc:
            text_parts.append(page.get_text("text"))
    return "\n".join(text_parts)


def _extract_docx_text(file_path: str) -> str:
    doc = Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())
    return "\n".join(paragraphs)


def _clean_text(text: str) -> str:
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    text = re.sub(r"[^\x20-\x7E\n]", " ", text)
    return text.strip()


def _extract_email(text: str) -> str:
    pattern = r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}"
    match = re.search(pattern, text)
    return match.group(0) if match else ""


def _extract_phone(text: str) -> str:
    pattern = r"(?:\+91[\-\s]?)?[6-9]\d{9}|(?:\+1[\-\s]?)?\(?\d{3}\)?[\-\s]?\d{3}[\-\s]?\d{4}"
    match = re.search(pattern, text)
    return match.group(0) if match else ""


def _extract_name(text: str) -> str:
    """
    Heuristic: the candidate name is usually on the first non-empty line
    and contains only letters and spaces (no @, digits, URLs).
    """
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    for line in lines[:5]:
        if (
            len(line.split()) <= 5
            and re.match(r"^[A-Za-z\s\.\-]+$", line)
            and len(line) > 3
        ):
            return line.title()
    return "Unknown Candidate"
